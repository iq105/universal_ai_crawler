"""导出服务：CSV / JSON / xlsx / docx / pdf"""
import csv
import io
import json
import logging

from openpyxl import Workbook

from app.services import result_store

logger = logging.getLogger(__name__)

try:  # python-docx
    from docx import Document
    from docx.shared import Pt
    _DOCX_OK = True
except Exception as exc:  # noqa: BLE001
    logger.warning("python-docx 加载失败，Word 导出将不可用：%s", exc)
    _DOCX_OK = False

try:  # reportlab
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    _PDF_OK = True
except Exception as exc:  # noqa: BLE001
    logger.warning("reportlab 加载失败，PDF 导出将不可用：%s", exc)
    _PDF_OK = False


def _union_keys(rows: list[dict]) -> list[str]:
    keys: list[str] = []
    seen = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                keys.append(k)
    return keys


async def export(task_id: str, fmt: str) -> tuple[bytes, str, str]:
    rows = await result_store.get_results(task_id, 0, 100000)
    data = [r["data"] for r in rows]
    filename = f"{task_id}.{fmt}"
    if fmt == "json":
        content = json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")
        return content, filename, "application/json; charset=utf-8"
    if fmt == "xlsx":
        cols = _union_keys(data)
        wb = Workbook()
        ws = wb.active
        ws.title = "数据"
        ws.append(cols)
        for d in data:
            ws.append([d.get(c, "") for c in cols])
        bio = io.BytesIO()
        wb.save(bio)
        return bio.getvalue(), filename, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if fmt == "docx":
        return _export_docx(data), filename, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if fmt == "pdf":
        return _export_pdf(data), filename, "application/pdf"
    # 默认 csv
    cols = _union_keys(data)
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=cols, extrasaction="ignore")
    writer.writeheader()
    for d in data:
        writer.writerow(d)
    return buf.getvalue().encode("utf-8-sig"), filename, "text/csv; charset=utf-8"


def _export_docx(rows: list[dict]) -> bytes:
    """导出为 Word 文档：标题 + 逐条记录表格"""
    if not _DOCX_OK:
        raise RuntimeError("python-docx 未安装，无法导出 Word")
    doc = Document()
    doc.add_heading("爬虫导出数据", level=1)
    doc.add_paragraph(f"共 {len(rows)} 条")
    for i, d in enumerate(rows, 1):
        doc.add_heading(f"记录 {i}", level=2)
        table = doc.add_table(rows=0, cols=2)
        table.style = "Light Grid Accent 1"
        for k, v in d.items():
            row_cells = table.add_row().cells
            row_cells[0].text = str(k)
            row_cells[0].paragraphs[0].runs[0].bold = True
            cell_text = str(v) if v is not None else ""
            row_cells[1].text = cell_text
        doc.add_paragraph("")
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def _export_pdf(rows: list[dict]) -> bytes:
    """导出为 PDF：landscape A4 表格"""
    if not _PDF_OK:
        raise RuntimeError("reportlab 未安装，无法导出 PDF")
    if not rows:
        # 空表也输出合法 PDF
        bio = io.BytesIO()
        doc = SimpleDocTemplate(bio, pagesize=A4)
        doc.build([Paragraph("无数据", getSampleStyleSheet()["Normal"])])
        return bio.getvalue()
    cols = _union_keys(rows)
    # 中文表格需要注册中文字体
    try:
        from reportlab.pdfbase.ttfonts import TTFont

        import os

        for name, path in (
            ("simsun", r"C:\Windows\Fonts\simsun.ttc"),
            ("msyh", r"C:\Windows\Fonts\msyh.ttc"),
            ("simhei", r"C:\Windows\Fonts\simhei.ttf"),
        ):
            if os.path.exists(path):
                pdfmetrics.registerFont(TTFont(name, path))
                break
    except Exception as exc:  # noqa: BLE001
        logger.warning("PDF 中文字体注册失败，将回退默认字体：%s", exc)

    font_name = "simsun" if "simsun" in pdfmetrics.getRegisteredFontNames() else "Helvetica"

    doc = SimpleDocTemplate(
        io.BytesIO(), pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
        topMargin=10 * mm, bottomMargin=10 * mm,
    )
    styles = getSampleStyleSheet()
    story = [Paragraph("爬虫导出数据", styles["Title"]), Spacer(1, 6 * mm)]

    table_data = [[Paragraph(f"<b>{c}</b>", styles["Normal"]) for c in cols]]
    for d in rows:
        table_data.append([Paragraph(_pdf_escape(str(d.get(c, ""))), styles["Normal"]) for c in cols])

    col_w = doc.width / max(len(cols), 1)
    # 限制单元格内容避免超长溢出：超长截断
    table = Table(table_data, repeatRows=1, colWidths=[col_w] * len(cols))
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#D9E2F3")]),
            ]
        )
    )
    story.append(table)
    bio = io.BytesIO()
    doc.build(story)
    return bio.getvalue()


def _pdf_escape(text: str, max_len: int = 200) -> str:
    if len(text) > max_len:
        text = text[:max_len] + "…"
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
